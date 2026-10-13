"""MySQL-backed coverage for daily leave V2 and template-backed documents."""

from __future__ import annotations

import hashlib
import io
import os
import secrets
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4
from zoneinfo import ZoneInfo

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import delete, func, select, text

from app.auth import get_current_member, hash_token, utcnow
from app.db import SessionLocal, get_db
from app.leave_documents import build_camp_leave_college_docx
from app.main import app
from app.models import (
    AdminUpdateEvent,
    AdminUpdateRead,
    CampLeaveEvent,
    CampLeaveParticipant,
    DailyLeaveDocumentRequest,
    DailyLeaveEntry,
    DailyLeaveWindow,
    Member,
    MemberIdentityAudit,
    SchoolLeaveRequest,
    SchoolLeaveRun,
)


TZ = ZoneInfo("Asia/Shanghai")


def local_now() -> datetime:
    return datetime.now(TZ).replace(tzinfo=None, second=0, microsecond=0)


def docx_text(content: bytes) -> tuple[list[str], list[list[str]]]:
    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    return paragraphs, tables


def package_media(content: bytes) -> set[bytes]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return {archive.read(name) for name in archive.namelist() if "/media/" in name or name.startswith("media/")}


def download_credential() -> str:
    return secrets.token_urlsafe(32)


def main() -> None:
    suffix = uuid4().hex
    admin_email = f"leave-v2-admin-{suffix}@example.invalid"
    second_admin_email = f"leave-v2-admin-second-{suffix}@example.invalid"
    formal_email = f"leave-v2-formal-{suffix}@example.invalid"
    reserve_email = f"leave-v2-reserve-{suffix}@example.invalid"
    unassigned_email = f"leave-v2-unassigned-{suffix}@example.invalid"
    incomplete_email = f"leave-v2-incomplete-{suffix}@example.invalid"
    event_ids: set[int] = set()
    window_ids: set[int] = set()
    legacy_run_ids: set[int] = set()
    previous_overrides = dict(app.dependency_overrides)
    original_phone = os.environ.get("LEAVE_CONTACT_PHONE")
    os.environ["LEAVE_CONTACT_PHONE"] = "000-0000-0000"

    with SessionLocal() as db:
        assert db.scalar(text("SELECT @@session.time_zone")) == "+00:00"
        admin = Member(name="V2 测试管理员", email=admin_email, role="admin", status="active")
        second_admin = Member(
            name="V2 测试管理员二", email=second_admin_email, role="admin", status="active"
        )
        formal = Member(
            name="正式成员甲",
            email=formal_email,
            student_id="26000201",
            college="electronic-science-engineering",
            team_membership="formal",
            role="member",
            status="active",
        )
        reserve = Member(
            name="Reserve Member",
            email=reserve_email,
            student_id="26000202",
            college="artificial-intelligence",
            team_membership="reserve",
            role="member",
            status="active",
        )
        unassigned = Member(
            name="身份待管理员维护",
            email=unassigned_email,
            student_id="26000204",
            college="art",
            team_membership=None,
            role="member",
            status="active",
        )
        incomplete = Member(
            name="资料未全",
            email=incomplete_email,
            student_id=None,
            college=None,
            team_membership="formal",
            role="member",
            status="active",
        )
        db.add_all([admin, second_admin, formal, reserve, unassigned, incomplete])
        db.commit()
        db.refresh(admin)
        db.refresh(second_admin)
        db.refresh(formal)
        db.refresh(reserve)
        db.refresh(unassigned)
        db.refresh(incomplete)

        def override_database():
            yield db

        app.dependency_overrides[get_db] = override_database
        actor: Member | None = None

        def set_actor(member: Member | None) -> None:
            nonlocal actor
            actor = member
            if member is None:
                app.dependency_overrides.pop(get_current_member, None)
            else:
                app.dependency_overrides[get_current_member] = lambda: actor

        client = TestClient(app)
        try:
            now = local_now()
            activity_start = (now + timedelta(days=3)).replace(hour=13, minute=30)
            activity_end = activity_start.replace(hour=17, minute=10)

            # Team members can generate directly without an administrator-created window.
            standalone_payload = {
                "start_at": activity_start.isoformat(timespec="minutes"),
                "end_at": activity_end.isoformat(timespec="minutes"),
            }
            set_actor(formal)
            standalone_doc = client.post("/api/daily-leave/self-service/document", json=standalone_payload)
            assert standalone_doc.status_code == 200, standalone_doc.text
            standalone_ready = standalone_doc.json()
            assert standalone_ready["download_url"].startswith("/api/daily-leave/self-service/document/")
            assert standalone_payload["start_at"] not in standalone_ready["download_url"]
            standalone_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id.is_(None), DailyLeaveEntry.member_id == formal.id
                ).order_by(DailyLeaveEntry.id)
            )
            assert standalone_entry is not None
            configured_phone = os.environ.pop("LEAVE_CONTACT_PHONE")
            try:
                failed_download = client.get(standalone_ready["download_url"])
                assert failed_download.status_code == 503
            finally:
                os.environ["LEAVE_CONTACT_PHONE"] = configured_phone
            assert db.scalar(
                select(func.count(DailyLeaveDocumentRequest.id)).where(
                    DailyLeaveDocumentRequest.entry_id == standalone_entry.id
                )
            ) == 2
            standalone_file = client.get(standalone_ready["download_url"])
            assert standalone_file.status_code == 200, standalone_file.text
            assert "attachment" in standalone_file.headers["content-disposition"].lower()
            assert standalone_file.headers["cache-control"].startswith("private, no-store")
            assert standalone_file.headers["x-content-type-options"] == "nosniff"
            standalone_paragraphs, standalone_tables = docx_text(standalone_file.content)
            assert any(
                "吉甲大师双创基地机器人战队创新实践活动" in line.replace("\u2060", "")
                for line in standalone_paragraphs
            )
            assert not any("内部活动" in line for line in standalone_paragraphs)
            assert any("下午课程" in line for line in standalone_paragraphs)
            assert standalone_tables[0][1] == ["正式成员甲", "26000201", "电子科学与工程学院"]
            assert standalone_entry.start_at == activity_start.replace(second=0, microsecond=0)
            assert standalone_entry.end_at == activity_end.replace(second=0, microsecond=0)
            assert standalone_entry.name_snapshot == "正式成员甲"
            assert standalone_entry.student_id_snapshot == "26000201"
            assert standalone_entry.college_snapshot == "electronic-science-engineering"
            assert standalone_entry.participant_type == "formal"
            assert standalone_entry.submitted_at is not None
            first_request_log = list(
                db.scalars(
                    select(DailyLeaveDocumentRequest)
                    .where(DailyLeaveDocumentRequest.entry_id == standalone_entry.id)
                    .order_by(DailyLeaveDocumentRequest.id)
                )
            )
            assert len(first_request_log) == 3
            assert first_request_log[0].request_type == "first_generation"
            assert first_request_log[0].request_stage == "prepare"
            assert first_request_log[1].request_stage == "download"
            assert first_request_log[1].request_type == "first_generation"
            assert first_request_log[2].request_stage == "download"
            assert first_request_log[2].request_type == "repeat_download"
            assert first_request_log[0].file_version == "daily-leave-v2"
            standalone_offline = client.post(
                "/api/daily-leave/self-service/document?offline=true", json=standalone_payload
            )
            assert standalone_offline.status_code == 200
            offline_ready = standalone_offline.json()
            assert offline_ready["filename"].startswith("线下签章版_")
            standalone_offline_file = client.get(offline_ready["download_url"])
            assert standalone_offline_file.status_code == 200
            assert package_media(standalone_offline_file.content) == set()
            request_logs = list(
                db.scalars(
                    select(DailyLeaveDocumentRequest)
                    .where(DailyLeaveDocumentRequest.entry_id == standalone_entry.id)
                    .order_by(DailyLeaveDocumentRequest.id)
                )
            )
            assert [item.request_stage for item in request_logs] == [
                "prepare", "download", "download", "prepare", "download"
            ]
            assert [item.request_type for item in request_logs] == [
                "first_generation", "first_generation", "repeat_download", "repeat_download", "repeat_download"
            ]
            assert [item.offline for item in request_logs] == [False, False, False, True, True]
            assert db.scalar(
                select(func.count(AdminUpdateEvent.id)).where(
                    AdminUpdateEvent.event_type == "leave_record_created",
                    AdminUpdateEvent.entity_type == "daily_leave_entry",
                    AdminUpdateEvent.entity_id == standalone_entry.id,
                )
            ) == 1
            assert client.get(standalone_ready["download_url"]).content == standalone_file.content
            assert db.scalar(
                select(func.count(AdminUpdateEvent.id)).where(
                    AdminUpdateEvent.entity_type == "daily_leave_entry",
                    AdminUpdateEvent.entity_id == standalone_entry.id,
                )
            ) == 1

            assert client.post(
                "/api/daily-leave/self-service/document",
                json={**standalone_payload, "title": "不允许客户端提交活动名"},
            ).status_code == 422
            assert client.post(
                "/api/daily-leave/self-service/document",
                json={**standalone_payload, "end_at": standalone_payload["start_at"]},
            ).status_code == 422
            set_actor(unassigned)
            unassigned_self_doc = client.post("/api/daily-leave/self-service/document", json=standalone_payload)
            assert unassigned_self_doc.status_code == 200, unassigned_self_doc.text
            unassigned_self_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id.is_(None), DailyLeaveEntry.member_id == unassigned.id
                )
            )
            assert unassigned_self_entry is not None and unassigned_self_entry.participant_type == "other"
            assert db.scalar(
                select(func.count(DailyLeaveEntry.id)).where(
                    DailyLeaveEntry.window_id.is_(None), DailyLeaveEntry.member_id == formal.id
                )
            ) == 1

            # Requests in parallel must serialize even when neither request sees a prior row.
            race_start = (now + timedelta(days=9)).replace(hour=8, minute=0)
            race_payload = {
                "start_at": race_start.isoformat(timespec="minutes"),
                "end_at": race_start.replace(hour=17, minute=10).isoformat(timespec="minutes"),
            }
            set_actor(formal)
            barrier = Barrier(2)

            def independent_request_db():
                with SessionLocal() as request_db:
                    yield request_db

            def concurrent_generation() -> int:
                with TestClient(app) as concurrent_client:
                    barrier.wait(timeout=10)
                    response = concurrent_client.post(
                        "/api/daily-leave/self-service/document", json=race_payload
                    )
                    return response.status_code

            db.commit()
            db.refresh(formal)
            previous_db_override = app.dependency_overrides[get_db]
            app.dependency_overrides[get_db] = independent_request_db
            try:
                with ThreadPoolExecutor(max_workers=2) as executor:
                    concurrent_statuses = list(executor.map(lambda _: concurrent_generation(), range(2)))
            finally:
                app.dependency_overrides[get_db] = previous_db_override
            db.commit()
            db.expire_all()
            assert concurrent_statuses == [200, 200]
            concurrent_entry_count = db.scalar(
                select(func.count(DailyLeaveEntry.id)).where(
                    DailyLeaveEntry.window_id.is_(None),
                    DailyLeaveEntry.member_id == formal.id,
                    DailyLeaveEntry.start_at == race_start,
                    DailyLeaveEntry.end_at == race_start.replace(hour=17, minute=10),
                )
            )
            assert concurrent_entry_count == 1, f"expected one parallel upsert row, found {concurrent_entry_count}"
            raced_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id.is_(None),
                    DailyLeaveEntry.member_id == formal.id,
                    DailyLeaveEntry.start_at == race_start,
                    DailyLeaveEntry.end_at == race_start.replace(hour=17, minute=10),
                )
            )
            assert raced_entry is not None
            raced_request_types = list(
                db.scalars(
                    select(DailyLeaveDocumentRequest.request_type)
                    .where(DailyLeaveDocumentRequest.entry_id == raced_entry.id)
                    .order_by(DailyLeaveDocumentRequest.id)
                )
            )
            assert sorted(raced_request_types) == ["first_generation", "repeat_download"]
            assert db.scalar(
                select(func.count(AdminUpdateEvent.id)).where(
                    AdminUpdateEvent.entity_type == "daily_leave_entry",
                    AdminUpdateEvent.entity_id == raced_entry.id,
                )
            ) == 1
            download_barrier = Barrier(2)

            def concurrent_file_request() -> int:
                with TestClient(app) as concurrent_client:
                    download_barrier.wait(timeout=10)
                    response = concurrent_client.get(
                        f"/api/daily-leave/self-service/document/{raced_entry.id}"
                    )
                    return response.status_code

            db.commit()
            previous_db_override = app.dependency_overrides[get_db]
            app.dependency_overrides[get_db] = independent_request_db
            try:
                with ThreadPoolExecutor(max_workers=2) as executor:
                    download_statuses = list(executor.map(lambda _: concurrent_file_request(), range(2)))
            finally:
                app.dependency_overrides[get_db] = previous_db_override
                db.expire_all()
            assert download_statuses == [200, 200]
            raced_download_types = list(
                db.scalars(
                    select(DailyLeaveDocumentRequest.request_type)
                    .where(
                        DailyLeaveDocumentRequest.entry_id == raced_entry.id,
                        DailyLeaveDocumentRequest.request_stage == "download",
                    )
                    .order_by(DailyLeaveDocumentRequest.id)
                )
            )
            assert raced_download_types == ["first_generation", "repeat_download"]
            assert db.scalar(
                select(func.count(DailyLeaveEntry.id)).where(
                    DailyLeaveEntry.window_id.is_(None),
                    DailyLeaveEntry.member_id == formal.id,
                    DailyLeaveEntry.start_at == race_start,
                    DailyLeaveEntry.end_at == race_start.replace(hour=17, minute=10),
                )
            ) == 1
            assert db.scalar(
                select(func.count(AdminUpdateEvent.id)).where(
                    AdminUpdateEvent.entity_type == "daily_leave_entry",
                    AdminUpdateEvent.entity_id == raced_entry.id,
                )
            ) == 1
            set_actor(reserve)
            request_count_before_forbidden_download = db.scalar(
                select(func.count(DailyLeaveDocumentRequest.id)).where(
                    DailyLeaveDocumentRequest.entry_id == standalone_entry.id
                )
            )
            assert client.get(standalone_ready["download_url"]).status_code == 404
            assert db.scalar(
                select(func.count(DailyLeaveDocumentRequest.id)).where(
                    DailyLeaveDocumentRequest.entry_id == standalone_entry.id
                )
            ) == request_count_before_forbidden_download
            set_actor(formal)
            set_actor(incomplete)
            assert client.post("/api/daily-leave/self-service/document", json=standalone_payload).status_code == 409

            window_payload = {
                "title": "管理员内部留档名称",
                "start_at": activity_start.isoformat(timespec="minutes"),
                "end_at": activity_end.isoformat(timespec="minutes"),
            }

            set_actor(formal)
            assert client.post("/api/daily-leave/admin/windows", json=window_payload).status_code == 403
            set_actor(admin)
            invalid_window = {**window_payload, "end_at": window_payload["start_at"]}
            assert client.post("/api/daily-leave/admin/windows", json=invalid_window).status_code == 422
            created = client.post("/api/daily-leave/admin/windows", json=window_payload)
            assert created.status_code == 201, created.text
            window = created.json()
            window_id = window["id"]
            window_ids.add(window_id)
            assert window["open_until"] == window["end_at"]
            assert window["public_enabled"] is False
            assert window["public_path"] is None
            assert window["team_open"] is True
            set_actor(formal)
            assert client.post(f"/api/daily-leave/admin/windows/{window_id}/public-link").status_code == 403
            set_actor(admin)
            enabled_link = client.post(f"/api/daily-leave/admin/windows/{window_id}/public-link")
            assert enabled_link.status_code == 200, enabled_link.text
            window = enabled_link.json()
            token = window["public_path"].rsplit("/", 1)[1]
            assert window["public_enabled"] is True
            assert 40 <= len(token) <= 50
            assert set(token) <= set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_-")

            # Formal and reserve members generate immediately from snapshots; an incomplete profile is rejected.
            set_actor(formal)
            member_windows = client.get("/api/daily-leave/windows")
            assert member_windows.status_code == 200
            assert member_windows.json()[0]["title"] == window_payload["title"]
            formal_doc = client.post(f"/api/daily-leave/windows/{window_id}/document")
            assert formal_doc.status_code == 200, formal_doc.text
            formal_ready = formal_doc.json()
            formal_file = client.get(formal_ready["download_url"])
            assert formal_file.status_code == 200, formal_file.text
            first_entry = db.scalar(select(DailyLeaveEntry).where(DailyLeaveEntry.window_id == window_id))
            assert first_entry is not None
            assert first_entry.participant_type == "formal"
            assert first_entry.name_snapshot == "正式成员甲"
            assert first_entry.student_id_snapshot == "26000201"
            assert first_entry.college_snapshot == "electronic-science-engineering"
            assert first_entry.start_at is None and first_entry.end_at is None
            set_actor(admin)
            assert client.get(formal_ready["download_url"]).status_code == 404
            set_actor(formal)
            formal_repeat = client.post(f"/api/daily-leave/windows/{window_id}/document")
            assert formal_repeat.status_code == 200
            formal_repeat_file = client.get(formal_repeat.json()["download_url"])
            assert formal_repeat_file.status_code == 200
            assert db.scalar(select(func.count(DailyLeaveEntry.id)).where(DailyLeaveEntry.window_id == window_id)) == 1

            set_actor(reserve)
            reserve_doc = client.post(f"/api/daily-leave/windows/{window_id}/document")
            assert reserve_doc.status_code == 200, reserve_doc.text
            reserve_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id == window_id,
                    DailyLeaveEntry.member_id == reserve.id,
                )
            )
            assert reserve_entry is not None and reserve_entry.participant_type == "reserve"

            set_actor(unassigned)
            unassigned_windows = client.get("/api/daily-leave/windows")
            assert unassigned_windows.status_code == 200
            assert any(item["id"] == window_id for item in unassigned_windows.json())
            unassigned_doc = client.post(f"/api/daily-leave/windows/{window_id}/document")
            assert unassigned_doc.status_code == 200, unassigned_doc.text
            unassigned_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id == window_id,
                    DailyLeaveEntry.member_id == unassigned.id,
                )
            )
            assert unassigned_entry is not None and unassigned_entry.participant_type == "other"

            set_actor(incomplete)
            assert client.post(f"/api/daily-leave/windows/{window_id}/document").status_code == 409

            # Default daily DOCX embeds exact source media and deterministic body/date/identity; offline omits both images.
            signed_bytes = formal_file.content
            signed_paragraphs, signed_tables = docx_text(signed_bytes)
            assert not any(window_payload["title"] in line for line in signed_paragraphs)
            time_text = (
                f"{activity_start.year}年{activity_start.month}月{activity_start.day}日"
                "13:30 至 17:10"
            )
            assert any(
                "吉甲大师双创基地机器人战队创新实践活动" in line.replace("\u2060", "")
                for line in signed_paragraphs
            )
            assert any("下午课程" in line for line in signed_paragraphs)
            assert any(time_text == line.removeprefix("以下学生因参加").split("的")[0] for line in signed_paragraphs if line.startswith("以下学生因参加"))
            assert signed_tables[0][0] == ["姓名", "学号", "学院"]
            assert signed_tables[0][1] == ["正式成员甲", "26000201", "电子科学与工程学院"]
            signature = (Path(__file__).parents[1] / "app/templates/assets/instructor-signature.jpg").read_bytes()
            seal = (Path(__file__).parents[1] / "app/templates/assets/team-seal.jpg").read_bytes()
            assert {hashlib.sha256(value).hexdigest() for value in package_media(signed_bytes)} >= {
                hashlib.sha256(signature).hexdigest(),
                hashlib.sha256(seal).hexdigest(),
            }
            set_actor(formal)
            offline = client.post(f"/api/daily-leave/windows/{window_id}/document?offline=true")
            assert offline.status_code == 200, offline.text
            offline_file = client.get(offline.json()["download_url"])
            assert offline_file.status_code == 200
            assert package_media(offline_file.content) == set()
            offline_paragraphs, offline_tables = docx_text(offline_file.content)
            assert [paragraph.replace("\u00a0", "") for paragraph in offline_paragraphs] == signed_paragraphs
            assert offline_tables == signed_tables

            # Public link exposes only fixed activity fields and repeats the original participant snapshot.
            set_actor(None)
            public_info = client.get(f"/api/daily-leave/public/{token}")
            assert public_info.status_code == 200
            public_json = public_info.json()
            assert public_json["start_at"].startswith(window_payload["start_at"])
            assert public_json["end_at"].startswith(window_payload["end_at"])
            assert "title" not in public_json and "open_until" not in public_json
            assert not {"id", "participants", "entries", "entry_count", "public_token", "public_path"}.intersection(public_json)
            external_payload = {
                "name": "临时参与者",
                "student_id": "26000203",
                "college": "computer-science-technology",
                "download_credential": download_credential(),
            }
            external_doc = client.post(f"/api/daily-leave/public/{token}/document", json=external_payload)
            assert external_doc.status_code == 200, external_doc.text
            external_ready = external_doc.json()
            # Model a POST that committed while its response was lost to a timed-out
            # browser. Retrying from the saved tab capability must reuse the same row.
            timeout_retry = client.post(f"/api/daily-leave/public/{token}/document", json=external_payload)
            assert timeout_retry.status_code == 200, timeout_retry.text
            assert timeout_retry.json()["download_url"] == external_ready["download_url"]
            assert db.scalar(
                select(func.count(DailyLeaveEntry.id)).where(
                    DailyLeaveEntry.window_id == window_id,
                    DailyLeaveEntry.student_id_snapshot == external_payload["student_id"],
                )
            ) == 1
            external_file = client.get(external_ready["download_url"])
            assert external_file.status_code == 200, external_file.text
            assert "attachment" in external_file.headers["content-disposition"].lower()
            assert external_file.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            assert external_file.headers["cache-control"] == "private, no-store, no-cache, max-age=0"
            assert external_file.headers["pragma"] == "no-cache"
            assert external_file.headers["x-content-type-options"] == "nosniff"
            assert external_file.headers["referrer-policy"] == "no-referrer"
            assert token not in external_ready["download_url"]
            assert external_payload["student_id"] not in external_ready["download_url"]
            assert external_payload["name"] not in external_ready["download_url"]
            assert "/public-download/" in external_ready["download_url"]
            external_repeat = client.post(
                f"/api/daily-leave/public/{token}/document",
                json={**external_payload, "name": "其他输入"},
            )
            assert external_repeat.status_code == 200
            assert external_repeat.json()["download_url"] == external_ready["download_url"]
            external_repeat_file = client.get(external_repeat.json()["download_url"])
            assert external_repeat_file.status_code == 200
            _, repeat_table = docx_text(external_repeat_file.content)
            assert repeat_table[0][1] == ["临时参与者", "26000203", "计算机科学与技术学院"]
            assert db.scalar(
                select(func.count(DailyLeaveEntry.id)).where(
                    DailyLeaveEntry.window_id == window_id,
                    DailyLeaveEntry.student_id_snapshot == external_payload["student_id"],
                )
            ) == 1
            public_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id == window_id,
                    DailyLeaveEntry.student_id_snapshot == external_payload["student_id"],
                )
            )
            assert public_entry is not None and public_entry.participant_type == "other"
            assert public_entry.public_download_token_hash == hash_token(external_payload["download_credential"])
            assert public_entry.public_download_expires_at is not None
            assert public_entry.public_download_expires_at > utcnow()
            public_request_logs = list(
                db.scalars(
                    select(DailyLeaveDocumentRequest)
                    .where(DailyLeaveDocumentRequest.entry_id == public_entry.id)
                    .order_by(DailyLeaveDocumentRequest.id)
                )
            )
            assert [item.request_stage for item in public_request_logs] == [
                "prepare", "prepare", "download", "prepare", "download"
            ]
            assert [item.request_type for item in public_request_logs] == [
                "first_generation",
                "repeat_download",
                "first_generation",
                "repeat_download",
                "repeat_download",
            ]
            assert db.scalar(
                select(func.count(AdminUpdateEvent.id)).where(
                    AdminUpdateEvent.entity_type == "daily_leave_entry",
                    AdminUpdateEvent.entity_id == public_entry.id,
                )
            ) == 1
            wrong_credential = {**external_payload, "download_credential": download_credential()}
            refused_claim = client.post(f"/api/daily-leave/public/{token}/document", json=wrong_credential)
            assert refused_claim.status_code == 409, f"unexpected status {refused_claim.status_code}"
            assert client.get(
                f"/api/daily-leave/public/{token}/document/{public_entry.id}"
            ).status_code == 404
            # Student number alone and a different random credential cannot recover an existing snapshot.
            assert client.post(
                f"/api/daily-leave/public/{token}/document",
                json={key: value for key, value in external_payload.items() if key != "download_credential"},
            ).status_code == 422
            public_paragraphs, public_tables = docx_text(external_file.content)
            assert public_tables[0][1] == ["临时参与者", "26000203", "计算机科学与技术学院"]
            assert "entries" not in "".join(public_paragraphs)
            assert client.post(
                f"/api/daily-leave/public/{token}/document",
                json={**external_payload, "student_id": "１２３４５６７８"},
            ).status_code == 422
            assert client.post(
                f"/api/daily-leave/public/{token}/document",
                json={**external_payload, "college": "invalid-college"},
            ).status_code == 422
            assert client.post(
                f"/api/daily-leave/public/{token}/document",
                json={**external_payload, "start_at": window_payload["start_at"]},
            ).status_code == 422

            original_capability_expiry = public_entry.public_download_expires_at
            public_entry.public_download_expires_at = utcnow() - timedelta(seconds=1)
            db.commit()
            expired_capability = client.get(external_ready["download_url"])
            assert expired_capability.status_code == 410
            assert expired_capability.headers["cache-control"].startswith("private, no-store")
            assert client.post(
                f"/api/daily-leave/public/{token}/document", json=external_payload
            ).status_code == 410
            public_entry.public_download_expires_at = original_capability_expiry
            db.commit()

            # Admin identity corrections affect new matters but preserve existing snapshots.
            set_actor(admin)
            identity_update = client.patch(
                f"/api/members/{formal.id}/profile",
                json={
                    "name": "后来改名",
                    "student_id": "26000991",
                    "college": "art",
                    "team_membership": "reserve",
                },
            )
            assert identity_update.status_code == 200, identity_update.text
            assert identity_update.json()["identity_locked"] is True
            set_actor(formal)
            old_self_file = client.get(standalone_ready["download_url"])
            assert old_self_file.status_code == 200
            _, old_self_tables = docx_text(old_self_file.content)
            assert old_self_tables[0][1] == ["正式成员甲", "26000201", "电子科学与工程学院"]
            new_start = (now + timedelta(days=11)).replace(hour=8, minute=0)
            new_self = client.post(
                "/api/daily-leave/self-service/document",
                json={
                    "start_at": new_start.isoformat(timespec="minutes"),
                    "end_at": new_start.replace(hour=17, minute=10).isoformat(timespec="minutes"),
                },
            )
            assert new_self.status_code == 200
            new_self_file = client.get(new_self.json()["download_url"])
            assert new_self_file.status_code == 200
            _, new_self_tables = docx_text(new_self_file.content)
            assert new_self_tables[0][1] == ["后来改名", "26000991", "艺术学院"]
            repeated_doc = client.post(f"/api/daily-leave/windows/{window_id}/document")
            assert repeated_doc.status_code == 200
            repeated_file = client.get(repeated_doc.json()["download_url"])
            assert repeated_file.status_code == 200
            repeated_paragraphs, repeated_tables = docx_text(repeated_file.content)
            assert repeated_tables[0][1] == signed_tables[0][1]
            assert db.scalar(select(func.count(DailyLeaveEntry.id)).where(DailyLeaveEntry.window_id == window_id)) == 4

            # Closing or reaching the deadline blocks both signed-in and public generation.
            set_actor(admin)
            closed = client.post(f"/api/daily-leave/admin/windows/{window_id}/close")
            assert closed.status_code == 200 and closed.json()["status"] == "closed"
            set_actor(formal)
            assert client.post(f"/api/daily-leave/windows/{window_id}/document").status_code == 409
            set_actor(None)
            # The holder of a still-valid capability may recover/re-download after closure.
            assert client.post(f"/api/daily-leave/public/{token}/document", json=external_payload).status_code == 200
            closed_public_info = client.get(f"/api/daily-leave/public/{token}")
            assert closed_public_info.status_code == 200
            assert closed_public_info.json()["accepting_participants"] is False
            assert client.get(external_ready["download_url"]).status_code == 200
            assert client.post(
                f"/api/daily-leave/public/{token}/document", json=external_payload
            ).status_code == 200
            assert client.post(
                f"/api/daily-leave/public/{token}/document",
                json={**external_payload, "student_id": "26000205"},
            ).status_code == 409
            assert client.get("/api/daily-leave/public/not-a-valid-token").status_code == 404

            set_actor(admin)
            expired_payload = {
                "title": "已截止的窗口",
                "start_at": (now + timedelta(days=5)).replace(hour=13, minute=30).isoformat(timespec="minutes"),
                "end_at": (now + timedelta(days=5)).replace(hour=17, minute=10).isoformat(timespec="minutes"),
            }
            expired_response = client.post("/api/daily-leave/admin/windows", json=expired_payload)
            assert expired_response.status_code == 201, expired_response.text
            expired_window = expired_response.json()
            window_ids.add(expired_window["id"])
            expired_row = db.get(DailyLeaveWindow, expired_window["id"])
            assert expired_row is not None
            expired_row.public_enabled = True
            expired_row.public_token = "expired-test-token-123456789012345678901234"
            expired_row.start_at = local_now() - timedelta(hours=2)
            expired_row.end_at = local_now() - timedelta(hours=1)
            expired_row.open_until = expired_row.end_at
            db.commit()
            expired_token = expired_row.public_token
            set_actor(None)
            expired_public_info = client.get(f"/api/daily-leave/public/{expired_token}")
            assert expired_public_info.status_code == 200
            assert expired_public_info.json()["accepting_participants"] is False
            assert client.post(
                f"/api/daily-leave/public/{expired_token}/document",
                json={
                    "name": "截止后报名",
                    "student_id": "26000205",
                    "college": "art",
                    "download_credential": download_credential(),
                },
            ).status_code == 409

            # Camp docs use immutable college snapshots, with one separate list for each college.
            set_actor(admin)
            camp_payload = {
                "title": "测试冬令营活动",
                "type": "winter",
                "start_date": (date.today() + timedelta(days=20)).isoformat(),
                "end_date": (date.today() + timedelta(days=25)).isoformat(),
                "collection_deadline": (local_now() + timedelta(days=3)).isoformat(timespec="minutes"),
            }
            camp_response = client.post("/api/camp-leave/admin/events", json=camp_payload)
            assert camp_response.status_code == 201, camp_response.text
            camp_id = camp_response.json()["id"]
            event_ids.add(camp_id)
            set_actor(reserve)
            assert client.post(f"/api/camp-leave/events/{camp_id}/join").status_code == 201
            set_actor(formal)
            db.commit()
            assert client.post(f"/api/camp-leave/events/{camp_id}/join").status_code == 201
            set_actor(admin)
            assert client.post(f"/api/camp-leave/admin/events/{camp_id}/close").status_code == 200
            signed_camp = client.get(
                f"/api/camp-leave/admin/events/{camp_id}/colleges/art/document"
            )
            assert signed_camp.status_code == 200, signed_camp.text
            camp_paragraphs, camp_tables = docx_text(signed_camp.content)
            visible_camp_paragraphs = [line.replace("\u2060", "") for line in camp_paragraphs]
            assert any("吉甲大师双创基地机器人战队寒假创新实践活动" in line for line in visible_camp_paragraphs)
            assert not any("测试冬令营活动" in line for line in camp_paragraphs)
            assert any("日常管理" in line and "全天候的实验室活动安排" in line for line in camp_paragraphs)
            assert camp_tables[0][0] == ["序号", "姓名", "学号", "学院"]
            assert len(camp_tables[0]) == 2
            assert camp_tables[0][1][1:] == ["后来改名", "26000991", "艺术学院"]
            assert {hashlib.sha256(value).hexdigest() for value in package_media(signed_camp.content)} >= {
                hashlib.sha256(signature).hexdigest(), hashlib.sha256(seal).hexdigest()
            }
            other_college = client.get(
                f"/api/camp-leave/admin/events/{camp_id}/colleges/artificial-intelligence/document"
            )
            assert other_college.status_code == 200
            _, other_table = docx_text(other_college.content)
            assert len(other_table[0]) == 2
            assert other_table[0][1][1:] == ["Reserve Member", "26000202", "人工智能学院"]
            member_doc = client.get(f"/api/camp-leave/events/{camp_id}/document")
            assert member_doc.status_code == 404  # Admin has not joined as a camp participant.
            set_actor(reserve)
            member_doc = client.get(f"/api/camp-leave/events/{camp_id}/document")
            assert member_doc.status_code == 200
            member_paragraphs, member_tables = docx_text(member_doc.content)
            assert len(member_tables[0]) == 2
            assert member_tables[0][1][1:] == ["Reserve Member", "26000202", "人工智能学院"]
            assert any("寒假" in line for line in member_paragraphs)
            set_actor(admin)
            offline_camp = client.get(
                f"/api/camp-leave/admin/events/{camp_id}/colleges/art/document?offline=true"
            )
            assert offline_camp.status_code == 200
            assert package_media(offline_camp.content) == set()
            offline_camp_paragraphs, offline_camp_tables = docx_text(offline_camp.content)
            assert [paragraph.replace("\u00a0", "") for paragraph in offline_camp_paragraphs] == camp_paragraphs
            assert offline_camp_tables == camp_tables
            summer_content = build_camp_leave_college_docx(
                title="测试夏令营活动",
                event_type="summer",
                start_date=date(2027, 7, 1),
                end_date=date(2027, 7, 8),
                participants=[("夏令营学生", "26000299", "人工智能学院")],
                offline=True,
            )
            summer_paragraphs, summer_tables = docx_text(summer_content)
            visible_summer_paragraphs = [line.replace("\u2060", "") for line in summer_paragraphs]
            assert any("吉甲大师双创基地机器人战队暑假创新实践活动" in line for line in visible_summer_paragraphs)
            assert not any("测试夏令营活动" in line for line in summer_paragraphs)
            assert summer_tables[0][1][1:] == ["夏令营学生", "26000299", "人工智能学院"]
            assert package_media(summer_content) == set()

            # Existing aggregation tables, status model, and read-only history endpoint remain usable.
            legacy_run = SchoolLeaveRun(
                collected_at=local_now(),
                created_by=admin.id,
                reason="legacy history fixture",
                status="ready",
            )
            db.add(legacy_run)
            db.flush()
            legacy_run_ids.add(legacy_run.id)
            legacy_request = SchoolLeaveRequest(
                member_id=admin.id,
                start_at=now + timedelta(days=1),
                end_at=now + timedelta(days=1, hours=2),
                member_name_snapshot="历史申请人",
                student_id_snapshot="26000290",
                status="included",
                run_id=legacy_run.id,
            )
            db.add(legacy_request)
            legacy_duplicate = DailyLeaveEntry(
                window_id=None,
                member_id=formal.id,
                start_at=standalone_entry.start_at,
                end_at=standalone_entry.end_at,
                name_snapshot=standalone_entry.name_snapshot,
                student_id_snapshot=standalone_entry.student_id_snapshot,
                college_snapshot=standalone_entry.college_snapshot,
                participant_type=standalone_entry.participant_type,
                submitted_at=standalone_entry.submitted_at + timedelta(minutes=1),
            )
            db.add(legacy_duplicate)

            # An interval ending exactly at the selected day's midnight does not overlap that day.
            boundary_start = datetime.combine(
                (now + timedelta(days=5)).date(), datetime.min.time()
            )
            boundary_window = DailyLeaveWindow(
                title="边界共享活动",
                start_at=boundary_start - timedelta(hours=2),
                end_at=boundary_start,
                open_until=boundary_start,
                team_open=True,
                public_enabled=False,
                status="closed",
                created_by=admin.id,
            )
            db.add(boundary_window)
            db.flush()
            window_ids.add(boundary_window.id)
            db.add_all(
                [
                    DailyLeaveEntry(
                        window_id=None,
                        member_id=formal.id,
                        start_at=boundary_start - timedelta(hours=1),
                        end_at=boundary_start,
                        name_snapshot="边界生成成员",
                        student_id_snapshot="26000231",
                        college_snapshot="electronic-science-engineering",
                        participant_type="formal",
                    ),
                    DailyLeaveEntry(
                        window_id=boundary_window.id,
                        member_id=formal.id,
                        start_at=None,
                        end_at=None,
                        name_snapshot="边界共享成员",
                        student_id_snapshot="26000233",
                        college_snapshot="electronic-science-engineering",
                        participant_type="formal",
                    ),
                    SchoolLeaveRequest(
                        member_id=formal.id,
                        start_at=boundary_start - timedelta(hours=1),
                        end_at=boundary_start,
                        member_name_snapshot="边界正式申请人",
                        student_id_snapshot="26000232",
                    ),
                ]
            )
            db.commit()
            db.refresh(legacy_request)
            set_actor(formal)
            assert client.get("/api/school-leave/admin/records").status_code == 403

            set_actor(admin)
            unified_records = client.get("/api/school-leave/admin/records")
            assert unified_records.status_code == 200, unified_records.text
            generated_rows = [
                row for row in unified_records.json()
                if row["name"] == "正式成员甲" and row["record_type"] == "generated"
            ]
            assert any(row["source_type"] == "self_service" for row in generated_rows)
            self_service_row = next(row for row in generated_rows if row["source_type"] == "self_service")
            assert self_service_row["id"] == f"generated:{legacy_duplicate.id}"
            assert self_service_row["generation_count"] == 2
            assert len(self_service_row["generation_times"]) == 2
            assert self_service_row["document_request_count"] == 7
            assert self_service_row["prepare_request_count"] == 2
            assert self_service_row["download_request_count"] == 5
            assert [item["request_stage"] for item in self_service_row["document_requests"]] == [
                "prepare", "download", "download", "prepare", "download", "download", "download"
            ]
            assert [item["request_type"] for item in self_service_row["document_requests"]] == [
                "first_generation",
                "first_generation",
                "repeat_download",
                "repeat_download",
                "repeat_download",
                "repeat_download",
                "repeat_download",
            ]
            assert {item["file_version"] for item in self_service_row["document_requests"]} == {
                "daily-leave-v2"
            }
            assert self_service_row["is_unread"] is True
            assert len(self_service_row["update_event_ids"]) == 1
            assert {row["source_type"] for row in generated_rows} == {"self_service", "shared_activity"}
            expected_generation_times = [
                value.replace(tzinfo=timezone.utc).astimezone(TZ).replace(tzinfo=None)
                for value in (standalone_entry.submitted_at, legacy_duplicate.submitted_at)
            ]
            assert [datetime.fromisoformat(value) for value in self_service_row["generation_times"]] == sorted(
                expected_generation_times
            )
            self_service_recorded_at = datetime.fromisoformat(self_service_row["recorded_at"])
            expected_self_service_time = (
                legacy_duplicate.submitted_at.replace(tzinfo=timezone.utc)
                .astimezone(TZ)
                .replace(tzinfo=None)
            )
            assert self_service_recorded_at == expected_self_service_time

            unread_before = client.get("/api/admin-updates/unread-count")
            assert unread_before.status_code == 200 and unread_before.json()["unread_count"] > 0
            marked_read = client.post(
                "/api/admin-updates/read",
                json={"event_ids": self_service_row["update_event_ids"]},
            )
            assert marked_read.status_code == 200, marked_read.text
            assert marked_read.json()["marked_count"] == 1
            admin_read = db.scalar(
                select(AdminUpdateRead).where(
                    AdminUpdateRead.event_id == self_service_row["update_event_ids"][0],
                    AdminUpdateRead.admin_id == admin.id,
                )
            )
            assert admin_read is not None
            reread_row = next(
                row for row in client.get("/api/school-leave/admin/records").json()
                if row["id"] == self_service_row["id"]
            )
            assert reread_row["is_unread"] is False

            set_actor(second_admin)
            second_admin_row = next(
                row for row in client.get(
                    "/api/school-leave/admin/records",
                    params={
                        "name": "正式成员甲",
                        "date_from": activity_start.date().isoformat(),
                        "record_type": "generated",
                    },
                ).json()
                if row["id"] == self_service_row["id"]
            )
            assert second_admin_row["is_unread"] is True
            assert db.scalar(
                select(AdminUpdateRead.id).where(
                    AdminUpdateRead.event_id == self_service_row["update_event_ids"][0],
                    AdminUpdateRead.admin_id == second_admin.id,
                )
            ) is None
            set_actor(formal)
            assert client.get("/api/admin-updates/unread-count").status_code == 403
            assert client.post(
                "/api/admin-updates/read",
                json={"event_ids": self_service_row["update_event_ids"]},
            ).status_code == 403

            leave_request_start = (now + timedelta(days=14)).replace(
                minute=0,
                second=0,
                microsecond=0,
            )
            created_leave_request = client.post(
                "/api/school-leave/requests",
                json={
                    "start_at": leave_request_start.isoformat(timespec="minutes"),
                    "end_at": (leave_request_start + timedelta(hours=2)).isoformat(timespec="minutes"),
                },
            )
            assert created_leave_request.status_code == 201, created_leave_request.text
            set_actor(admin)
            new_application_record = client.get(
                "/api/school-leave/admin/records",
                params={
                    "name": "后来改名",
                    "date_from": leave_request_start.date().isoformat(),
                    "record_type": "formal_application",
                },
            )
            assert new_application_record.status_code == 200
            new_application = new_application_record.json()[0]
            assert new_application["is_unread"] is True
            assert len(new_application["update_event_ids"]) == 1

            legacy_application_rows = client.get(
                "/api/school-leave/admin/records",
                params={
                    "name": "历史申请人",
                    "date_from": (now + timedelta(days=1)).date().isoformat(),
                    "record_type": "formal_application",
                },
            )
            assert legacy_application_rows.status_code == 200
            legacy_application = legacy_application_rows.json()[0]
            assert legacy_application["is_unread"] is False
            assert legacy_application["update_event_ids"] == []

            admin_unread_before_delete = client.get("/api/admin-updates/unread-count").json()["unread_count"]
            set_actor(second_admin)
            second_admin_unread_before_delete = client.get("/api/admin-updates/unread-count").json()["unread_count"]
            set_actor(admin)
            deletable_run = SchoolLeaveRun(
                collected_at=local_now(),
                created_by=admin.id,
                reason="leave request deletion audit fixture",
                status="completed",
            )
            db.add(deletable_run)
            db.flush()
            legacy_run_ids.add(deletable_run.id)
            deletable_request = SchoolLeaveRequest(
                member_id=formal.id,
                start_at=leave_request_start + timedelta(days=1),
                end_at=leave_request_start + timedelta(days=1, hours=2),
                member_name_snapshot="待删除历史申请",
                student_id_snapshot="26000201",
                status="included",
                run_id=deletable_run.id,
            )
            db.add(deletable_request)
            db.flush()
            deletable_event = AdminUpdateEvent(
                event_type="leave_record_created",
                entity_type="school_leave_request",
                entity_id=deletable_request.id,
            )
            db.add(deletable_event)
            db.flush()
            deletable_read = AdminUpdateRead(event_id=deletable_event.id, admin_id=admin.id)
            db.add(deletable_read)
            db.commit()
            deletable_run_id = deletable_run.id
            deletable_request_id = deletable_request.id
            deletable_event_id = deletable_event.id
            deletable_read_id = deletable_read.id
            assert client.get("/api/admin-updates/unread-count").json()["unread_count"] == admin_unread_before_delete
            set_actor(second_admin)
            assert client.get("/api/admin-updates/unread-count").json()["unread_count"] == second_admin_unread_before_delete + 1
            set_actor(admin)
            deleted_history = client.delete(f"/api/school-leave/admin/runs/{deletable_run_id}")
            assert deleted_history.status_code == 204, deleted_history.text
            db.expire_all()
            assert db.get(SchoolLeaveRequest, deletable_request_id) is None
            assert db.get(AdminUpdateEvent, deletable_event_id) is None
            assert db.get(AdminUpdateRead, deletable_read_id) is None
            assert client.get("/api/admin-updates/unread-count").json()["unread_count"] == admin_unread_before_delete
            set_actor(second_admin)
            assert client.get("/api/admin-updates/unread-count").json()["unread_count"] == second_admin_unread_before_delete
            set_actor(admin)

            shared_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id == window_id,
                    DailyLeaveEntry.member_id == formal.id,
                )
            )
            assert shared_entry is not None
            shared_row = next(
                row for row in unified_records.json()
                if row["source_type"] == "shared_activity"
                and row["activity_title"] == window_payload["title"]
                and row["student_id"] == "26000201"
            )
            assert shared_row["start_at"].startswith(window_payload["start_at"])
            assert shared_row["record_type"] == "generated"
            assert shared_row["application_status"] is None
            shared_recorded_at = datetime.fromisoformat(shared_row["recorded_at"])
            expected_shared_time = (
                shared_entry.submitted_at.replace(tzinfo=timezone.utc)
                .astimezone(TZ)
                .replace(tzinfo=None)
            )
            assert shared_recorded_at == expected_shared_time

            legacy_records = client.get(
                "/api/school-leave/admin/records",
                params={
                    "name": "历史申请人",
                    "date_from": (now + timedelta(days=1)).date().isoformat(),
                    "date_to": (now + timedelta(days=1)).date().isoformat(),
                    "record_type": "formal_application",
                },
            )
            assert legacy_records.status_code == 200, legacy_records.text
            assert len(legacy_records.json()) == 1
            legacy_record = legacy_records.json()[0]
            assert legacy_record["record_type"] == "formal_application"
            assert legacy_record["source_type"] == "legacy_application"
            assert legacy_record["application_status"] == "included"
            assert legacy_record["is_unread"] is False
            assert legacy_record["activity_title"] == "legacy history fixture"
            assert legacy_record["activity_at"] is not None
            legacy_recorded_at = datetime.fromisoformat(legacy_record["recorded_at"])
            expected_legacy_time = (
                legacy_request.created_at.replace(tzinfo=timezone.utc)
                .astimezone(TZ)
                .replace(tzinfo=None)
            )
            assert legacy_recorded_at == expected_legacy_time

            for boundary_name, boundary_type in (
                ("边界生成成员", "generated"),
                ("边界共享成员", "generated"),
                ("边界正式申请人", "formal_application"),
            ):
                boundary_response = client.get(
                    "/api/school-leave/admin/records",
                    params={
                        "name": boundary_name,
                        "date_from": boundary_start.date().isoformat(),
                        "record_type": boundary_type,
                    },
                )
                assert boundary_response.status_code == 200, boundary_response.text
                assert boundary_response.json() == []

            assert client.get(
                "/api/school-leave/admin/records",
                params={"date_from": "2026-10-10", "date_to": "2026-10-09"},
            ).status_code == 422

            history_list = client.get("/api/school-leave/admin/runs")
            assert history_list.status_code == 200
            assert any(run["id"] == legacy_run.id for run in history_list.json())
            history_doc = client.get(f"/api/school-leave/admin/runs/{legacy_run.id}/history-document")
            assert history_doc.status_code == 200, history_doc.text
            assert Document(io.BytesIO(history_doc.content)).tables
        finally:
            app.dependency_overrides.clear()
            app.dependency_overrides.update(previous_overrides)
            fixture_member_ids = [admin.id, second_admin.id, formal.id, reserve.id, unassigned.id, incomplete.id]
            generated_ids = set(
                db.scalars(
                    select(DailyLeaveEntry.id).where(DailyLeaveEntry.member_id.in_(fixture_member_ids))
                )
            )
            if window_ids:
                generated_ids.update(
                    db.scalars(
                        select(DailyLeaveEntry.id).where(DailyLeaveEntry.window_id.in_(window_ids))
                    )
                )
            application_ids = set(
                db.scalars(
                    select(SchoolLeaveRequest.id).where(SchoolLeaveRequest.member_id.in_(fixture_member_ids))
                )
            )
            if generated_ids or application_ids:
                from sqlalchemy import or_

                event_filters = []
                if generated_ids:
                    event_filters.append(
                        (AdminUpdateEvent.entity_type == "daily_leave_entry")
                        & AdminUpdateEvent.entity_id.in_(generated_ids)
                    )
                if application_ids:
                    event_filters.append(
                        (AdminUpdateEvent.entity_type == "school_leave_request")
                        & AdminUpdateEvent.entity_id.in_(application_ids)
                    )
                db.execute(delete(AdminUpdateEvent).where(or_(*event_filters)))
            if window_ids:
                db.execute(delete(DailyLeaveEntry).where(DailyLeaveEntry.window_id.in_(window_ids)))
                db.execute(delete(DailyLeaveWindow).where(DailyLeaveWindow.id.in_(window_ids)))
            db.execute(
                delete(DailyLeaveEntry).where(
                    DailyLeaveEntry.member_id.in_(fixture_member_ids)
                )
            )
            if event_ids:
                db.execute(delete(CampLeaveParticipant).where(CampLeaveParticipant.event_id.in_(event_ids)))
                db.execute(delete(CampLeaveEvent).where(CampLeaveEvent.id.in_(event_ids)))
            if legacy_run_ids:
                db.execute(
                    delete(SchoolLeaveRequest).where(
                        SchoolLeaveRequest.member_id.in_(
                            fixture_member_ids
                        )
                    )
                )
                db.execute(delete(SchoolLeaveRun).where(SchoolLeaveRun.id.in_(legacy_run_ids)))
            db.execute(
                delete(MemberIdentityAudit).where(
                    MemberIdentityAudit.member_id.in_(fixture_member_ids)
                )
            )
            db.execute(delete(Member).where(Member.email.in_([admin_email, second_admin_email, formal_email, reserve_email, unassigned_email, incomplete_email])))
            db.commit()
            if original_phone is None:
                os.environ.pop("LEAVE_CONTACT_PHONE", None)
            else:
                os.environ["LEAVE_CONTACT_PHONE"] = original_phone

    print("Leave V2 workflow and DOCX tests passed")


if __name__ == "__main__":
    main()
